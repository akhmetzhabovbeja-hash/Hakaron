from docx import Document
doc = Document("/tmp/task.docx")
for p in doc.paragraphs:
    print(p.text)
for table in doc.tables:
    for row in table.rows:
        print(" | ".join(cell.text for cell in row.cells))
    print("---")
