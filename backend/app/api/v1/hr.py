import io
import os
import logging
from datetime import datetime, timezone, date, time

import httpx
from pydantic import BaseModel as PydanticBaseModel
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select, or_, func
from sqlalchemy.ext.asyncio import AsyncSession
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.api.v1.auth import get_current_user
from app.core.database import get_db
from app.models.user import User, UserRole
from app.models.vacancy import Vacancy
from app.models.question import Question, QuestionCategory
from app.models.vacancy_question import VacancyQuestion
from app.models.candidate import CandidateProfile
from app.models.analysis import CandidateAnalysis, AnalysisStatus
from app.schemas.vacancy import VacancyCreate, VacancyUpdate, VacancyResponse, VacancyDetailResponse, QuestionInVacancy
from app.schemas.question import QuestionResponse, QuestionCreate, VacancyQuestionAssign
from app.models.questionnaire import QuestionnaireResponse
from app.schemas.candidate import CandidateListResponse, CandidateAnalysisResponse, CandidateDossierResponse, AnswerItem

router = APIRouter()


def _require_hr(user: User) -> None:
    """Raise 403 if the user is not an HR."""
    if user.role != UserRole.HR:
        raise HTTPException(status_code=403, detail="HR role required")


def _build_candidate_list(analysis, profile, vacancy) -> CandidateListResponse:
    flags = analysis.ai_detection_flags or []
    ai_count = len([f for f in flags if f.get("is_ai_generated")])
    return CandidateListResponse(
        id=analysis.id,
        candidate_id=profile.id,
        full_name=profile.full_name,
        email=profile.email,
        vacancy_title=vacancy.title,
        total_score=analysis.total_score,
        vacancy_match=analysis.vacancy_match,
        status=analysis.status.value,
        source=profile.source.value,
        ai_flags_count=len(flags),
        ai_suspected=ai_count > 0,
    )


# ---------------------------------------------------------------------------
# Vacancies
# ---------------------------------------------------------------------------


@router.get("/vacancies", response_model=list[VacancyResponse])
async def list_vacancies(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List all vacancies (HR only)."""
    _require_hr(current_user)
    result = await db.execute(select(Vacancy).order_by(Vacancy.created_at.desc()))
    return result.scalars().all()


@router.post("/vacancies", response_model=VacancyResponse, status_code=201)
async def create_vacancy(
    data: VacancyCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Create a new vacancy (HR only)."""
    _require_hr(current_user)
    vacancy = Vacancy(
        title=data.title,
        description=data.description,
        requirements=data.requirements,
        application_deadline=data.application_deadline,
        created_by=current_user.id,
    )
    db.add(vacancy)
    await db.commit()
    await db.refresh(vacancy)
    return vacancy


@router.get("/vacancies/{vacancy_id}", response_model=VacancyDetailResponse)
async def get_vacancy(
    vacancy_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get vacancy details with assigned questions (HR only)."""
    _require_hr(current_user)

    result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")

    # Fetch questions joined through VacancyQuestion, ordered by VacancyQuestion.order
    q_result = await db.execute(
        select(Question, VacancyQuestion.order)
        .join(VacancyQuestion, VacancyQuestion.question_id == Question.id)
        .where(VacancyQuestion.vacancy_id == vacancy_id)
        .order_by(VacancyQuestion.order)
    )
    rows = q_result.all()

    questions = [
        QuestionInVacancy(
            id=question.id,
            text=question.text,
            category=question.category.value,
            order=order,
        )
        for question, order in rows
    ]

    return VacancyDetailResponse(
        id=vacancy.id,
        title=vacancy.title,
        description=vacancy.description,
        requirements=vacancy.requirements,
        is_active=vacancy.is_active,
        application_deadline=vacancy.application_deadline,
        questions=questions,
    )


@router.put("/vacancies/{vacancy_id}", response_model=VacancyResponse)
async def update_vacancy(
    vacancy_id: int,
    data: VacancyUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Update vacancy fields (HR only)."""
    _require_hr(current_user)
    result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")

    for field in ("title", "description", "requirements", "application_deadline"):
        value = getattr(data, field, None)
        if value is not None:
            setattr(vacancy, field, value)

    await db.commit()
    await db.refresh(vacancy)
    return vacancy


@router.patch("/vacancies/{vacancy_id}/archive", response_model=VacancyResponse)
async def archive_vacancy(
    vacancy_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Soft-delete: set is_active=False (HR only)."""
    _require_hr(current_user)
    result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")
    vacancy.is_active = False
    await db.commit()
    await db.refresh(vacancy)
    return vacancy


@router.patch("/vacancies/{vacancy_id}/restore", response_model=VacancyResponse)
async def restore_vacancy(
    vacancy_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Restore archived vacancy (HR only)."""
    _require_hr(current_user)
    result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")
    vacancy.is_active = True
    await db.commit()
    await db.refresh(vacancy)
    return vacancy


# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------


@router.get("/statistics")
async def get_statistics(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Aggregated statistics for HR dashboard."""
    _require_hr(current_user)

    # All analyses
    result = await db.execute(
        select(CandidateAnalysis, Vacancy)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
    )
    rows = result.all()

    if not rows:
        return {
            "total_candidates": 0, "avg_score": 0, "avg_match": 0,
            "by_status": {}, "by_vacancy": [], "score_distribution": [],
            "ai_flags_count": 0,
        }

    analyses = [a for a, _ in rows]

    total = len(analyses)
    avg_score = round(sum(a.total_score for a in analyses) / total, 1)
    avg_match = round(sum(a.vacancy_match for a in analyses) / total, 2)

    # By status
    by_status: dict[str, int] = {}
    for a in analyses:
        by_status[a.status.value] = by_status.get(a.status.value, 0) + 1

    # By vacancy
    vac_map: dict[int, dict] = {}
    for a, v in rows:
        if v.id not in vac_map:
            vac_map[v.id] = {"title": v.title, "count": 0, "total_score": 0}
        vac_map[v.id]["count"] += 1
        vac_map[v.id]["total_score"] += a.total_score
    by_vacancy = [
        {"title": d["title"], "count": d["count"], "avg_score": round(d["total_score"] / d["count"], 1)}
        for d in vac_map.values()
    ]

    # Score distribution
    ranges = [(90, 100), (80, 89), (70, 79), (60, 69), (0, 59)]
    score_distribution = []
    for lo, hi in ranges:
        cnt = sum(1 for a in analyses if lo <= a.total_score <= hi)
        score_distribution.append({"range": f"{lo}-{hi}", "count": cnt})

    # AI flags
    ai_flags_count = sum(1 for a in analyses if a.ai_detection_flags and len(a.ai_detection_flags) > 0)

    return {
        "total_candidates": total,
        "avg_score": avg_score,
        "avg_match": avg_match,
        "by_status": by_status,
        "by_vacancy": by_vacancy,
        "score_distribution": score_distribution,
        "ai_flags_count": ai_flags_count,
    }


# ---------------------------------------------------------------------------
# Questions bank
# ---------------------------------------------------------------------------


@router.get("/questions/bank", response_model=list[QuestionResponse])
async def list_questions(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List all system questions and custom questions created by this HR."""
    _require_hr(current_user)
    result = await db.execute(
        select(Question).where(
            or_(
                Question.is_system == True,  # noqa: E712
                Question.created_by == current_user.id,
            )
        )
    )
    return result.scalars().all()


@router.post("/questions", response_model=QuestionResponse, status_code=201)
async def create_question(
    data: QuestionCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Create a custom question (HR only)."""
    _require_hr(current_user)
    question = Question(
        text=data.text,
        category=QuestionCategory(data.category),
        is_system=False,
        created_by=current_user.id,
    )
    db.add(question)
    await db.commit()
    await db.refresh(question)
    return question


# ---------------------------------------------------------------------------
# Assign questions to vacancy
# ---------------------------------------------------------------------------


@router.post("/vacancies/{vacancy_id}/questions", status_code=201)
async def assign_questions_to_vacancy(
    vacancy_id: int,
    data: VacancyQuestionAssign,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Assign a list of questions to a vacancy (HR only).

    Replaces any previously assigned questions. Order is determined by list position.
    """
    _require_hr(current_user)

    # Verify vacancy exists
    result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")

    # Remove existing assignments
    existing = await db.execute(
        select(VacancyQuestion).where(VacancyQuestion.vacancy_id == vacancy_id)
    )
    for vq in existing.scalars().all():
        await db.delete(vq)

    # Create new assignments with order based on list position
    for idx, question_id in enumerate(data.question_ids):
        # Verify question exists
        q_result = await db.execute(select(Question).where(Question.id == question_id))
        if not q_result.scalar_one_or_none():
            raise HTTPException(
                status_code=404, detail=f"Question with id {question_id} not found"
            )
        vq = VacancyQuestion(
            vacancy_id=vacancy_id,
            question_id=question_id,
            order=idx,
        )
        db.add(vq)

    await db.commit()
    return {"message": "Questions assigned", "count": len(data.question_ids)}


# ---------------------------------------------------------------------------
# Candidates review
# ---------------------------------------------------------------------------


@router.get("/candidates", response_model=list[CandidateListResponse])
async def list_analyzed_candidates(
    search: str | None = Query(None),
    status_filter: str | None = Query(None),
    sort_by: str = Query("score"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List candidates with ANALYZED status (ready for HR review)."""
    _require_hr(current_user)

    stmt = (
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.status == AnalysisStatus.ANALYZED)
    )

    if search:
        stmt = stmt.where(
            CandidateProfile.full_name.ilike(f"%{search}%")
            | CandidateProfile.email.ilike(f"%{search}%")
        )

    if sort_by == "score":
        stmt = stmt.order_by(CandidateAnalysis.total_score.desc())
    elif sort_by == "date":
        stmt = stmt.order_by(CandidateAnalysis.created_at.desc())
    elif sort_by == "name":
        stmt = stmt.order_by(CandidateProfile.full_name)
    else:
        stmt = stmt.order_by(CandidateAnalysis.total_score.desc())

    result = await db.execute(stmt)
    rows = result.all()

    return [
        _build_candidate_list(analysis, profile, vacancy)
        for analysis, profile, vacancy in rows
    ]


@router.get("/candidates/{analysis_id}", response_model=CandidateAnalysisResponse)
async def get_candidate_analysis(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get full analysis details for HR review."""
    _require_hr(current_user)
    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.id == analysis_id)
    )
    row = result.one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Analysis not found")

    analysis, profile, vacancy = row

    return CandidateAnalysisResponse(
        id=analysis.id,
        candidate_id=profile.id,
        full_name=profile.full_name,
        email=profile.email,
        vacancy_title=vacancy.title,
        total_score=analysis.total_score,
        vacancy_match=analysis.vacancy_match,
        growth_potential=analysis.growth_potential,
        growth_path_score=analysis.growth_path_score,
        strengths=analysis.strengths or [],
        weaknesses=analysis.weaknesses or [],
        summary=analysis.summary,
        status=analysis.status.value,
        ai_detection_flags=analysis.ai_detection_flags or [],
        category_scores=analysis.category_scores,
        manager_comment=analysis.manager_comment,
    )


@router.post("/candidates/{analysis_id}/send-to-manager")
async def send_to_manager(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Send candidate analysis to manager for approval."""
    _require_hr(current_user)
    result = await db.execute(
        select(CandidateAnalysis).where(CandidateAnalysis.id == analysis_id)
    )
    analysis = result.scalar_one_or_none()
    if not analysis:
        raise HTTPException(status_code=404, detail="Analysis not found")

    if analysis.status not in (AnalysisStatus.ANALYZED, AnalysisStatus.HR_REVIEW):
        raise HTTPException(
            status_code=400,
            detail=f"Cannot send to manager from status '{analysis.status.value}'",
        )

    analysis.status = AnalysisStatus.SENT_TO_MANAGER
    analysis.hr_reviewed_by = current_user.id
    analysis.sent_to_manager_at = datetime.now(timezone.utc)
    await db.commit()

    return {"message": "Candidate sent to manager for review", "analysis_id": analysis_id}


@router.post("/candidates/{analysis_id}/reject")
async def hr_reject_candidate(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """HR rejects candidate (does not pass to commission)."""
    _require_hr(current_user)
    result = await db.execute(
        select(CandidateAnalysis).where(CandidateAnalysis.id == analysis_id)
    )
    analysis = result.scalar_one_or_none()
    if not analysis:
        raise HTTPException(status_code=404, detail="Analysis not found")

    analysis.status = AnalysisStatus.REJECTED
    analysis.hr_reviewed_by = current_user.id
    await db.commit()

    return {"message": "Candidate rejected by HR", "analysis_id": analysis_id}


# ---------------------------------------------------------------------------
# Approved candidates
# ---------------------------------------------------------------------------


@router.get("/approved", response_model=list[CandidateListResponse])
async def list_approved_candidates(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List candidates approved by manager."""
    _require_hr(current_user)
    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.status == AnalysisStatus.APPROVED)
        .order_by(CandidateAnalysis.total_score.desc())
    )
    rows = result.all()

    return [
        _build_candidate_list(analysis, profile, vacancy)
        for analysis, profile, vacancy in rows
    ]


# ---------------------------------------------------------------------------
# Invite candidate (placeholder)
# ---------------------------------------------------------------------------


@router.post("/candidates/{analysis_id}/invite")
async def invite_candidate(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Invite an approved candidate (placeholder for email integration)."""
    _require_hr(current_user)

    result = await db.execute(
        select(CandidateAnalysis).where(CandidateAnalysis.id == analysis_id)
    )
    analysis = result.scalar_one_or_none()
    if not analysis:
        raise HTTPException(status_code=404, detail="Analysis not found")

    if analysis.status != AnalysisStatus.APPROVED:
        raise HTTPException(
            status_code=400,
            detail="Only approved candidates can be invited",
        )

    return {"message": "Invitation sent successfully", "analysis_id": analysis_id}


# ---------------------------------------------------------------------------
# PDF / Excel font & label helpers
# ---------------------------------------------------------------------------

_fonts_registered = False


def _register_fonts():
    global _fonts_registered
    if _fonts_registered:
        return
    for path, name in [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "DejaVu"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "DejaVuBold"),
    ]:
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont(name, path))
    _fonts_registered = True


STATUS_LABELS = {
    "pending": "Ожидание",
    "processing": "Анализ",
    "analyzed": "Проанализирован",
    "hr_review": "На рассмотрении",
    "sent_to_manager": "У комиссии",
    "approved": "Зачислен",
    "rejected": "Не прошёл",
}


# ---------------------------------------------------------------------------
# Candidates per vacancy (all statuses)
# ---------------------------------------------------------------------------


@router.get("/vacancies/{vacancy_id}/candidates", response_model=list[CandidateListResponse])
async def list_vacancy_candidates(
    vacancy_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List ALL candidates for a specific vacancy (all statuses)."""
    _require_hr(current_user)

    # Verify vacancy exists
    vac_result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    if not vac_result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Vacancy not found")

    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.vacancy_id == vacancy_id)
        .order_by(CandidateAnalysis.total_score.desc())
    )
    rows = result.all()

    return [
        _build_candidate_list(analysis, profile, vacancy)
        for analysis, profile, vacancy in rows
    ]


# ---------------------------------------------------------------------------
# Candidate dossier
# ---------------------------------------------------------------------------


@router.get("/candidates/{analysis_id}/dossier", response_model=CandidateDossierResponse)
async def get_candidate_dossier(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get full dossier for a candidate including user info and questionnaire answers."""
    _require_hr(current_user)

    # Get analysis + profile + vacancy
    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.id == analysis_id)
    )
    row = result.one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Analysis not found")

    analysis, profile, vacancy = row

    # Get user info
    user_result = await db.execute(select(User).where(User.id == profile.user_id))
    user = user_result.scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Get questionnaire answers
    answers_result = await db.execute(
        select(QuestionnaireResponse)
        .where(QuestionnaireResponse.candidate_id == profile.id)
        .order_by(QuestionnaireResponse.question_number)
    )
    answers = answers_result.scalars().all()

    return CandidateDossierResponse(
        name=user.name,
        email=user.email,
        phone=user.phone,
        bio=user.bio,
        avatar_url=user.avatar_url,
        id_document_url=profile.id_document_url,
        total_score=analysis.total_score,
        vacancy_match=analysis.vacancy_match,
        growth_potential=analysis.growth_potential,
        growth_path_score=analysis.growth_path_score,
        strengths=analysis.strengths or [],
        weaknesses=analysis.weaknesses or [],
        summary=analysis.summary,
        status=analysis.status.value,
        vacancy_title=vacancy.title,
        ai_detection_flags=analysis.ai_detection_flags or [],
        category_scores=analysis.category_scores,
        manager_comment=analysis.manager_comment,
        answers=[
            AnswerItem(
                question_number=a.question_number,
                question_text=a.question_text,
                answer_text=a.answer_text,
            )
            for a in answers
        ],
    )


# ---------------------------------------------------------------------------
# PDF report for a single candidate
# ---------------------------------------------------------------------------


@router.get("/candidates/{analysis_id}/report-pdf")
async def get_candidate_report_pdf(
    analysis_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Generate a PDF dossier report for a candidate."""
    _require_hr(current_user)
    _register_fonts()

    # ---- data fetching (same as dossier) ----
    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile, Vacancy)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .join(Vacancy, CandidateAnalysis.vacancy_id == Vacancy.id)
        .where(CandidateAnalysis.id == analysis_id)
    )
    row = result.one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Analysis not found")

    analysis, profile, vacancy = row

    user_result = await db.execute(select(User).where(User.id == profile.user_id))
    user = user_result.scalar_one_or_none()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    answers_result = await db.execute(
        select(QuestionnaireResponse)
        .where(QuestionnaireResponse.candidate_id == profile.id)
        .order_by(QuestionnaireResponse.question_number)
    )
    answers = answers_result.scalars().all()

    # ---- build PDF ----
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm)

    base_font = "DejaVu" if os.path.exists("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf") else "Helvetica"
    bold_font = "DejaVuBold" if os.path.exists("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf") else "Helvetica-Bold"

    title_style = ParagraphStyle("title", fontName=bold_font, fontSize=18,
                                  textColor=HexColor("#2563eb"), spaceAfter=12)
    heading_style = ParagraphStyle("heading", fontName=bold_font, fontSize=13,
                                    textColor=HexColor("#1e3a5f"), spaceBefore=14, spaceAfter=6)
    normal_style = ParagraphStyle("normal", fontName=base_font, fontSize=10, leading=14)
    small_style = ParagraphStyle("small", fontName=base_font, fontSize=9, leading=12)
    bullet_style = ParagraphStyle("bullet", fontName=base_font, fontSize=10, leading=14,
                                   leftIndent=12, bulletIndent=0, bulletFontName=base_font)

    elements: list = []

    # 1. Title
    elements.append(Paragraph("inVision U — Досье абитуриента", title_style))
    elements.append(Spacer(1, 6 * mm))

    # 2. Profile block
    profile_data = [
        ["ФИО", user.name or "—"],
        ["Email", user.email or "—"],
        ["Телефон", user.phone or "—"],
        ["Программа", vacancy.title or "—"],
    ]
    profile_table = Table(profile_data, colWidths=[45 * mm, 120 * mm])
    profile_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), bold_font),
        ("FONTNAME", (1, 0), (1, -1), base_font),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    elements.append(profile_table)
    elements.append(Spacer(1, 4 * mm))

    # 3. AI Assessment block
    elements.append(Paragraph("AI-оценка", heading_style))
    score_headers = [
        Paragraph("<b>Общий балл</b>", small_style),
        Paragraph("<b>Соответствие</b>", small_style),
        Paragraph("<b>Потенциал</b>", small_style),
        Paragraph("<b>Траектория роста</b>", small_style),
    ]
    score_values = [
        str(analysis.total_score or 0),
        f"{round((analysis.vacancy_match or 0) * 100)}%",
        str(analysis.growth_potential or "—"),
        f"{round((analysis.growth_path_score or 0) * 100)}%",
    ]
    score_table = Table([score_headers, score_values], colWidths=[40 * mm] * 4)
    score_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), base_font),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#cbd5e1")),
        ("BACKGROUND", (0, 0), (-1, 0), HexColor("#eff6ff")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(score_table)
    elements.append(Spacer(1, 4 * mm))

    # 4. Strengths
    strengths = analysis.strengths or []
    if strengths:
        elements.append(Paragraph("Сильные стороны", heading_style))
        for s in strengths:
            elements.append(Paragraph(f"• {s}", bullet_style))

    # 5. Weaknesses
    weaknesses = analysis.weaknesses or []
    if weaknesses:
        elements.append(Paragraph("Зоны развития", heading_style))
        for w in weaknesses:
            elements.append(Paragraph(f"• {w}", bullet_style))

    # 6. AI Summary
    if analysis.summary:
        elements.append(Paragraph("AI-резюме", heading_style))
        elements.append(Paragraph(analysis.summary, normal_style))

    # 7. AI Detection
    flags = analysis.ai_detection_flags or []
    if flags:
        elements.append(Paragraph("AI-детекция", heading_style))
        for flag in flags:
            if isinstance(flag, dict):
                q_num = flag.get("question_number", "?")
                reason = flag.get("reason", str(flag))
                elements.append(Paragraph(f"• Вопрос {q_num}: {reason}", bullet_style))
            else:
                elements.append(Paragraph(f"• {flag}", bullet_style))

    # 8. ID Document
    if profile.id_document_url:
        elements.append(Paragraph("Удостоверение личности", heading_style))
        id_path = f"/app{profile.id_document_url}"
        if os.path.exists(id_path):
            from reportlab.platypus import Image
            try:
                img = Image(id_path, width=80 * mm, height=50 * mm)
                img.hAlign = "LEFT"
                elements.append(img)
            except Exception:
                elements.append(Paragraph(f"Файл: {profile.id_document_url}", small_style))
        else:
            elements.append(Paragraph(f"Файл загружен: {profile.id_document_url}", small_style))
        elements.append(Spacer(1, 4 * mm))

    # 9. Answers
    if answers:
        elements.append(Paragraph("Ответы на вопросы", heading_style))
        for a in answers:
            elements.append(Spacer(1, 2 * mm))
            elements.append(Paragraph(
                f"<b>Вопрос {a.question_number}:</b> {a.question_text}", normal_style
            ))
            elements.append(Paragraph(a.answer_text, small_style))

    doc.build(elements)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="report_{analysis_id}.pdf"',
        },
    )


# ---------------------------------------------------------------------------
# Excel report for all candidates in a vacancy
# ---------------------------------------------------------------------------


@router.get("/vacancies/{vacancy_id}/report-excel")
async def get_vacancy_report_excel(
    vacancy_id: int,
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Generate an Excel report for all candidates in a vacancy/program."""
    _require_hr(current_user)

    # Verify vacancy
    vac_result = await db.execute(select(Vacancy).where(Vacancy.id == vacancy_id))
    vacancy = vac_result.scalar_one_or_none()
    if not vacancy:
        raise HTTPException(status_code=404, detail="Vacancy not found")

    # Build query
    query = (
        select(CandidateAnalysis, CandidateProfile)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .where(CandidateAnalysis.vacancy_id == vacancy_id)
    )

    parsed_from = None
    parsed_to = None
    if date_from:
        parsed_from = date.fromisoformat(date_from)
        query = query.where(CandidateAnalysis.created_at >= datetime.combine(parsed_from, time.min, tzinfo=timezone.utc))
    if date_to:
        parsed_to = date.fromisoformat(date_to)
        query = query.where(CandidateAnalysis.created_at <= datetime.combine(parsed_to, time.max, tzinfo=timezone.utc))

    query = query.order_by(CandidateAnalysis.total_score.desc())
    result = await db.execute(query)
    rows = result.all()

    # ---- Build workbook ----
    wb = Workbook()

    # ----- Sheet 1: Абитуриенты -----
    ws1 = wb.active
    ws1.title = "Абитуриенты"

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563eb", end_color="2563eb", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["№", "ФИО", "Email", "Общий балл", "Соответствие%", "Потенциал",
               "Траектория роста%", "Статус", "AI-флаги", "Дата подачи"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, (analysis, profile) in enumerate(rows, 2):
        flags = analysis.ai_detection_flags or []
        flag_count = len(flags)
        status_label = STATUS_LABELS.get(analysis.status.value, analysis.status.value)
        created = analysis.created_at.strftime("%Y-%m-%d") if analysis.created_at else "—"

        ws1.cell(row=row_idx, column=1, value=row_idx - 1)
        ws1.cell(row=row_idx, column=2, value=profile.full_name or "—")
        ws1.cell(row=row_idx, column=3, value=profile.email or "—")
        ws1.cell(row=row_idx, column=4, value=analysis.total_score or 0)
        ws1.cell(row=row_idx, column=5, value=round((analysis.vacancy_match or 0) * 100, 1))
        ws1.cell(row=row_idx, column=6, value=analysis.growth_potential or "—")
        ws1.cell(row=row_idx, column=7, value=round((analysis.growth_path_score or 0) * 100, 1))
        ws1.cell(row=row_idx, column=8, value=status_label)
        ws1.cell(row=row_idx, column=9, value=flag_count)
        ws1.cell(row=row_idx, column=10, value=created)

    # Auto-width
    for col in ws1.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws1.column_dimensions[col_letter].width = min(max_len + 4, 50)

    # ----- Sheet 2: Сводка -----
    ws2 = wb.create_sheet("Сводка")
    bold_font = Font(bold=True)

    period_str = "Все время"
    if parsed_from and parsed_to:
        period_str = f"{parsed_from.isoformat()} — {parsed_to.isoformat()}"
    elif parsed_from:
        period_str = f"с {parsed_from.isoformat()}"
    elif parsed_to:
        period_str = f"по {parsed_to.isoformat()}"

    ws2.cell(row=1, column=1, value="Программа:").font = bold_font
    ws2.cell(row=1, column=2, value=vacancy.title)
    ws2.cell(row=2, column=1, value="Период:").font = bold_font
    ws2.cell(row=2, column=2, value=period_str)

    total_count = len(rows)
    avg_score = sum((a.total_score or 0) for a, _ in rows) / total_count if total_count else 0
    avg_match = sum((a.vacancy_match or 0) for a, _ in rows) / total_count if total_count else 0

    ws2.cell(row=4, column=1, value="Всего абитуриентов:").font = bold_font
    ws2.cell(row=4, column=2, value=total_count)
    ws2.cell(row=5, column=1, value="Средний балл:").font = bold_font
    ws2.cell(row=5, column=2, value=round(avg_score, 1))
    ws2.cell(row=6, column=1, value="Среднее соответствие:").font = bold_font
    ws2.cell(row=6, column=2, value=f"{round(avg_match * 100, 1)}%")

    ws2.cell(row=8, column=1, value="Распределение по статусам:").font = bold_font

    status_counts: dict[str, int] = {}
    ai_flagged = 0
    for analysis, _ in rows:
        st = analysis.status.value
        status_counts[st] = status_counts.get(st, 0) + 1
        if analysis.ai_detection_flags:
            ai_flagged += 1

    current_row = 9
    for status_key, count in status_counts.items():
        label = STATUS_LABELS.get(status_key, status_key)
        ws2.cell(row=current_row, column=1, value=f"{label}:")
        ws2.cell(row=current_row, column=2, value=count)
        current_row += 1

    ws2.cell(row=current_row, column=1, value="С AI-флагами:").font = bold_font
    ws2.cell(row=current_row, column=2, value=ai_flagged)

    # Auto-width for sheet 2
    for col in ws2.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws2.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="report_{vacancy_id}.xlsx"',
        },
    )


# ---------------------------------------------------------------------------
# AI Detection (SlopTotal integration)
# ---------------------------------------------------------------------------

SLOPTOTAL_URL = "http://ai-detection:8000"


class AiDetectRequest(PydanticBaseModel):
    text: str


class SendMessageRequest(PydanticBaseModel):
    message: str


@router.post("/candidates/{analysis_id}/send-message")
async def send_message_to_candidate(
    analysis_id: int,
    body: SendMessageRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Send a message/notification to candidate from HR."""
    _require_hr(current_user)

    result = await db.execute(
        select(CandidateAnalysis, CandidateProfile)
        .join(CandidateProfile, CandidateAnalysis.candidate_id == CandidateProfile.id)
        .where(CandidateAnalysis.id == analysis_id)
    )
    row = result.one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="Not found")

    analysis, profile = row
    if not profile.user_id:
        raise HTTPException(status_code=400, detail="Candidate has no account")

    from app.models.notification import Notification
    notif = Notification(
        user_id=profile.user_id,
        from_user_id=current_user.id,
        title="Сообщение от координатора отбора",
        message=body.message,
    )
    db.add(notif)
    await db.commit()

    # Also send via Telegram if linked
    user = (await db.execute(select(User).where(User.id == profile.user_id))).scalar_one_or_none()
    if user and user.telegram_id:
        try:
            bot_token = "8657277109:AAGPSvKgPIhRd6yd2MwDlmRAtksXSfguHN8"
            async with httpx.AsyncClient(timeout=5.0) as client:
                await client.post(
                    f"https://api.telegram.org/bot{bot_token}/sendMessage",
                    json={"chat_id": user.telegram_id, "text": f"📩 Сообщение от координатора отбора:\n\n{body.message}"},
                )
        except Exception:
            pass

    return {"success": True}


@router.post("/ai-detect")
async def ai_detect_text(
    body: AiDetectRequest,
    current_user: User = Depends(get_current_user),
):
    """Hybrid AI detection: SlopTotal ML + Russian heuristics (HR only)."""
    _require_hr(current_user)

    if len(body.text.strip()) < 50:
        raise HTTPException(status_code=400, detail="Текст должен быть не менее 50 символов")

    from app.api.v1.candidates import detect_ai_text_async
    result = await detect_ai_text_async(body.text)

    # Verdict mapping
    prob = result.get("ai_probability", 0)
    if prob >= 0.7:
        verdict = "AI-текст"
    elif prob >= 0.41:
        verdict = "Подозрительный"
    else:
        verdict = "Человеческий"

    return {
        "source": "hybrid_ml+ru",
        "score": round(prob * 100, 1),
        "verdict": verdict,
        "ai_probability": prob,
        "is_ai_generated": result.get("is_ai_generated", False),
        "indicators": result.get("indicators", []),
        "ml_scores": result.get("ml_scores"),
        "russian_rule_score": result.get("russian_rule_score", 0),
        "sloptotal_engines": result.get("sloptotal_engines", []),
    }


@router.get("/ai-detect/status")
async def ai_detect_status(
    current_user: User = Depends(get_current_user),
):
    """Check SlopTotal service health."""
    _require_hr(current_user)
    try:
        async with httpx.AsyncClient(timeout=5.0) as client:
            resp = await client.get(f"{SLOPTOTAL_URL}/health")
            if resp.status_code == 200:
                return {"status": "online", "details": resp.json()}
    except Exception:
        pass
    return {"status": "offline"}
